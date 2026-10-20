import sys,json,os,uuid,html,subprocess,shutil,zipfile
from pathlib import Path
sys.path.insert(0,str(Path(__file__).resolve().parent))
ROOT=Path(os.environ.get('OPENCLAUDE_DOCUMENT_ROOT','/workspace')).resolve()
request=json.load(sys.stdin);op=request['operation'];a=request['args']
def scoped(value):
    p=(ROOT/value).resolve()
    if not p.is_relative_to(ROOT): raise ValueError('Path escaped workspace')
    if p.is_file() and p.stat().st_size>50*1024*1024: raise ValueError('Document input exceeds 50 MiB')
    if p.is_file() and p.suffix.lower() in ('.docx','.xlsx','.pptx'):
        with zipfile.ZipFile(p) as archive:
            entries=archive.infolist()
            if len(entries)>20000 or sum(item.file_size for item in entries)>256*1024*1024: raise ValueError('Expanded Office document exceeds safe processing limits')
    return p
def output(name,extension):
    name=Path(name).name
    if not name.lower().endswith('.'+extension): name+='.'+extension
    p=scoped('outputs/'+name);p.parent.mkdir(parents=True,exist_ok=True)
    if p.exists(): p=p.with_name(p.stem+'-'+uuid.uuid4().hex[:6]+p.suffix)
    return p
def record(path,mime): return {'path':str(path.relative_to(ROOT)),'mediaType':mime}
def preview(path):
    files=[]
    try:
        directory=scoped('previews/'+uuid.uuid4().hex);directory.mkdir(parents=True)
        if path.suffix.lower()=='.pdf': pdf=path
        else:
            office=shutil.which('libreoffice') or shutil.which('soffice')
            if not office and os.name=='nt':
                for base in (os.environ.get('ProgramFiles',r'C:\Program Files'),os.environ.get('ProgramFiles(x86)',r'C:\Program Files (x86)')):
                    candidate=Path(base)/'LibreOffice/program/soffice.exe'
                    if candidate.is_file(): office=str(candidate);break
            if not office: return [],'The file was created successfully. Office thumbnail conversion is unavailable because LibreOffice is not installed; Docker is not required.'
            profile=directory/'office-profile'
            subprocess.run([office,'-env:UserInstallation='+profile.as_uri(),'--headless','--convert-to','pdf','--outdir',str(directory),str(path)],capture_output=True,check=True,timeout=45)
            pdf=directory/(path.stem+'.pdf')
        if not pdf.exists(): return [],'Preview converter did not create a PDF'
        image=directory/'page.png'
        try:
            import pypdfium2 as pdfium
            document=pdfium.PdfDocument(str(pdf))
            try:
                page=document[0]
                try:
                    bitmap=page.render(scale=min(2,1400/max(page.get_width(),page.get_height())))
                    try: bitmap.to_pil().save(image)
                    finally: bitmap.close()
                finally: page.close()
            finally: document.close()
        except ImportError:
            subprocess.run(['pdftoppm','-f','1','-singlefile','-scale-to','1400','-png',str(pdf),str(directory/'page')],capture_output=True,check=True,timeout=20)
        if image.exists(): files.append(record(image,'image/png'))
        return files,None
    except Exception as error: return [],str(error)
result={}
if op=='document_extract':
    p=scoped(a['file_path']);ext=p.suffix.lower();text=''
    if ext=='.pdf':
        import pypdf
        reader=pypdf.PdfReader(p);text='\n\n'.join((page.extract_text() or '') for page in list(reader.pages)[:100])
        if not text.strip(): raise ValueError('No extractable text; scanned PDFs require an OCR adapter')
    elif ext=='.docx':
        from docx import Document
        document=Document(p);text='\n'.join(paragraph.text for paragraph in document.paragraphs)+'\n'+'\n'.join(' | '.join(cell.text for cell in row.cells) for table in document.tables for row in table.rows)
    elif ext=='.pptx':
        from pptx import Presentation
        text='\n'.join(shape.text for slide in Presentation(p).slides for shape in slide.shapes if hasattr(shape,'text'))
    elif ext=='.xlsx':
        import openpyxl
        book=openpyxl.load_workbook(p,read_only=True);text='\n'.join(str(row) for sheet in book for row in list(sheet.values)[:100]);book.close()
    else: text=p.read_text(errors='replace')
    result={'text':text[:30000],'truncated':len(text)>30000}
elif op=='document_create':
    fmt=a['format'];p=output(a['output_name'],fmt)
    if fmt=='docx':
        from docx import Document
        document=Document();document.add_heading(a['title'],0)
        for block in a['blocks']:
            if block['type']=='heading': document.add_heading(block.get('text',''),1)
            elif block['type']=='table':
                rows=block.get('rows',[])
                if rows:
                    table=document.add_table(rows=0,cols=max(len(row) for row in rows));table.style='Table Grid'
                    for row in rows:
                        cells=table.add_row().cells
                        for i,value in enumerate(row): cells[i].text=str(value)
            else: document.add_paragraph(block.get('text',''))
        document.save(p);Document(p)
        mime='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
    else:
        from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,Table
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        import reportlab
        fonts=[Path('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf'),Path(os.environ.get('WINDIR',r'C:\Windows'))/'Fonts/arial.ttf',Path(reportlab.__file__).resolve().parent/'fonts/Vera.ttf']
        font=next((path for path in fonts if path.is_file()),None)
        if font: pdfmetrics.registerFont(TTFont('OpenClaudeDocument',str(font)))
        styles=getSampleStyleSheet()
        for style in styles.byName.values(): style.fontName='OpenClaudeDocument' if font else 'Helvetica'
        story=[Paragraph(html.escape(a['title']),styles['Title']),Spacer(1,18)]
        for block in a['blocks']:
            if block['type']=='table': story.append(Table([[Paragraph(html.escape(str(v)),styles['BodyText']) for v in row] for row in block.get('rows',[])],hAlign='LEFT'))
            else: story.append(Paragraph(html.escape(block.get('text','')).replace('\n','<br/>'),styles['Heading2' if block['type']=='heading' else 'BodyText']))
            story.append(Spacer(1,12))
        SimpleDocTemplate(str(p)).build(story)
        import pypdf;pypdf.PdfReader(p);mime='application/pdf'
    previews,error=preview(p);result={'files':[record(p,mime),*previews],'preview_error':error}
elif op=='spreadsheet_create':
    import openpyxl
    p=output(a['output_name'],'xlsx');book=openpyxl.Workbook();book.remove(book.active)
    for source in a['sheets']:
        sheet=book.create_sheet(source['name'][:31])
        for row in source['rows']: sheet.append(row)
        sheet.freeze_panes='A2'
        for cell in sheet[1]: cell.font=openpyxl.styles.Font(bold=True);cell.fill=openpyxl.styles.PatternFill('solid',fgColor='EAE8E1')
        for column in sheet.columns: sheet.column_dimensions[column[0].column_letter].width=min(50,max(12,max(len(str(cell.value or '')) for cell in column)+2))
    book.save(p);check=openpyxl.load_workbook(p);check.close();previews,error=preview(p);result={'files':[record(p,'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'),*previews],'formulas_recalculated':False,'preview_error':error}
elif op=='spreadsheet_inspect':
    from document_editing import inspect_spreadsheet
    result=inspect_spreadsheet(scoped(a['file_path']),a)
elif op=='spreadsheet_edit':
    from document_editing import edit_spreadsheet
    p=output(a['output_name'],'xlsx');result=edit_spreadsheet(scoped(a['file_path']),p,a)
    previews,error=preview(p);result.update({'files':[record(p,'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'),*previews],'preview_error':error})
elif op=='pdf_inspect':
    from document_editing import inspect_pdf
    result=inspect_pdf(scoped(a['file_path']))
elif op=='pdf_edit':
    from document_editing import edit_pdf
    p=output(a['output_name'],'pdf');result=edit_pdf(scoped(a['file_path']),p,a)
    previews,error=preview(p);result.update({'files':[record(p,'application/pdf'),*previews],'preview_error':error})
elif op=='slides_create':
    from document_editing import create_slides
    p=output(a['output_name'],'pptx');create_slides(p,a)
    previews,error=preview(p);result={'files':[record(p,'application/vnd.openxmlformats-officedocument.presentationml.presentation'),*previews],'preview_error':error}
elif op=='slides_inspect':
    from document_editing import inspect_slides
    result=inspect_slides(scoped(a['file_path']))
elif op=='slides_edit':
    from document_editing import edit_slides
    p=output(a['output_name'],'pptx');result=edit_slides(scoped(a['file_path']),p,a)
    previews,error=preview(p);result.update({'files':[record(p,'application/vnd.openxmlformats-officedocument.presentationml.presentation'),*previews],'preview_error':error})
elif op=='spreadsheet_chart':
    from document_editing import chart_spreadsheet
    p=output(a['output_name'],'xlsx');result=chart_spreadsheet(scoped(a['file_path']),p,a)
    previews,error=preview(p);result.update({'files':[record(p,'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'),*previews],'preview_error':error})
elif op=='chart_create':
    import pandas as pd
    import matplotlib;matplotlib.use('Agg');import matplotlib.pyplot as plt
    data=pd.read_csv(scoped(a['file_path']));x=a['x'];y=a['y']
    if x not in data.columns or y not in data.columns: raise ValueError('Chart columns were not found')
    p=output('chart-'+uuid.uuid4().hex[:8],'png');fig,ax=plt.subplots(figsize=(10,6));getattr(ax,a['chart_type'] if a['chart_type']!='line' else 'plot')(data[x],data[y],color='#C76D4D');ax.set_title(a.get('title','Chart'));ax.set_xlabel(x);ax.set_ylabel(y);fig.tight_layout();fig.savefig(p,dpi=150);plt.close(fig);result={'files':[record(p,'image/png')],'rows_plotted':len(data)}
elif op=='image_transform':
    from PIL import Image
    image=Image.open(scoped(a['file_path']));operation=a['operation']
    if operation=='resize': image=image.resize((a['width'],a['height']))
    elif operation=='crop': image=image.crop(a['box'])
    elif operation=='rotate': image=image.rotate(a.get('angle',0),expand=True)
    fmt=a.get('format','png');p=output('image-'+uuid.uuid4().hex[:8],fmt)
    if fmt=='jpeg': image=image.convert('RGB')
    image.save(p);Image.open(p).verify();result={'files':[record(p,'image/'+fmt)]}
else: raise ValueError('Unknown document operation')
print(json.dumps(result,default=str))
