"""Editing helpers for the isolated document worker."""
from copy import copy
from io import BytesIO
from itertools import islice


def inspect_pdf(path):
    from pypdf import PdfReader
    reader = PdfReader(path)
    if reader.is_encrypted:
        raise ValueError('Encrypted PDFs must be decrypted before inspection or editing')
    fields = reader.get_fields() or {}
    return {
        'page_count': len(reader.pages),
        'pages': [{'page': index + 1, 'width': float(page.mediabox.width),
                   'height': float(page.mediabox.height), 'rotation': page.rotation}
                  for index, page in enumerate(islice(reader.pages, 1000))],
        'pages_truncated': len(reader.pages) > 1000,
        'fields': [{'name': name, 'type': str(field.get('/FT', '')), 'value': str(field.get('/V', ''))}
                   for name, field in fields.items()],
        'metadata': {str(key): str(value) for key, value in (reader.metadata or {}).items()},
    }


def edit_pdf(source, target, args):
    from pypdf import PdfReader, PdfWriter
    from reportlab.pdfgen.canvas import Canvas
    reader = PdfReader(source)
    if reader.is_encrypted:
        raise ValueError('Encrypted PDFs must be decrypted before editing')
    fields = reader.get_fields() or {}
    if any(field.get('/FT') == '/Sig' for field in fields.values()):
        raise ValueError('Editing signed PDFs is unsupported because it can invalidate signatures')
    selection = args.get('pages')
    if selection is not None and (not selection or any(type(page) is not int or page < 1 or page > len(reader.pages) for page in selection)):
        raise ValueError('Selected page numbers must be within the source PDF (1-based)')
    if selection is not None and fields:
        raise ValueError('Page selection on interactive forms is unsupported; omit pages to preserve fields')
    writer = PdfWriter()
    if selection is None:
        writer.clone_document_from_reader(reader)
    else:
        for number in selection:
            # Repeated source pages need independent content streams for later edits.
            writer.reset_translation(reader)
            writer.add_page(reader.pages[number - 1])
    def page_at(number):
        if type(number) is not int or number < 1 or number > len(writer.pages):
            raise ValueError('Edit page numbers refer to output pages (1-based)')
        return writer.pages[number - 1]
    for rotation in args.get('rotations', []):
        degrees = rotation['degrees']
        if type(degrees) is not int or degrees % 90:
            raise ValueError('Page rotation must be a multiple of 90 degrees')
        page_at(rotation['page']).rotate(degrees)
    for overlay in args.get('overlays', []):
        page = page_at(overlay['page'])
        stream = BytesIO()
        canvas = Canvas(stream, pagesize=(float(page.mediabox.width), float(page.mediabox.height)))
        canvas.setFont('Helvetica', overlay.get('font_size', 12))
        canvas.drawString(overlay['x'], overlay['y'], overlay['text'])
        canvas.save(); stream.seek(0)
        page.merge_page(PdfReader(stream).pages[0])
    values = args.get('form_fields', {})
    if values:
        missing = set(values) - set(fields)
        if missing:
            raise ValueError('Unknown form fields: ' + ', '.join(sorted(missing)))
        for name in values:
            if fields[name].get('/FT') != '/Tx':
                raise ValueError('Only text form fields are supported: ' + name)
        writer.update_page_form_field_values(None, values, auto_regenerate=False)
    with open(target, 'wb') as stream:
        writer.write(stream)
    actual = PdfReader(target)
    for name, value in values.items():
        if (actual.get_fields() or {}).get(name, {}).get('/V') != value:
            target.unlink(missing_ok=True)
            raise ValueError('PDF field verification failed: ' + name)
    return {'page_count': len(actual.pages), 'filled_fields': list(values),
            'overlay_notice': 'Text overlays add visible content; they do not remove or redact existing content.'}


def inspect_spreadsheet(path, args):
    import openpyxl
    book = openpyxl.load_workbook(path, read_only=True, data_only=False)
    try:
        if args.get('sheet') and args['sheet'] not in book.sheetnames:
            raise ValueError('Requested sheet was not found')
        max_rows = args.get('max_rows', 100); max_columns = args.get('max_columns', 50)
        sheets = []
        for sheet in book:
            if args.get('sheet') and sheet.title != args['sheet']:
                continue
            rows = [[value if isinstance(value, (str, int, float, bool, type(None))) else str(value)
                     for value in row] for row in sheet.iter_rows(max_row=min(sheet.max_row, max_rows),
                                                                max_col=min(sheet.max_column, max_columns), values_only=True)]
            sheets.append({'name': sheet.title, 'rows': rows, 'row_count': sheet.max_row,
                           'column_count': sheet.max_column,
                           'truncated': sheet.max_row > max_rows or sheet.max_column > max_columns})
        return {'sheets': sheets, 'truncated': any(sheet['truncated'] for sheet in sheets),
                'formulas_recalculated': False}
    finally:
        book.close()


def edit_spreadsheet(source, target, args):
    import openpyxl
    from openpyxl.utils.cell import range_boundaries, coordinate_to_tuple
    book = openpyxl.load_workbook(source)
    try:
        for change in args['sheets']:
            name = change['name']
            if name not in book.sheetnames:
                if not change.get('create', False):
                    raise ValueError('Requested sheet was not found: ' + name)
                book.create_sheet(name)
            sheet = book[name]
            for update in change.get('cells', []):
                row, column = coordinate_to_tuple(update['cell'])
                if row > 1048576 or column > 16384:
                    raise ValueError('Cell exceeds XLSX limits')
                sheet[update['cell']] = update['value']
            for fmt in change.get('formats', []):
                left, top, right, bottom = range_boundaries(fmt['range'])
                if not all((left, top, right, bottom)) or right < left or bottom < top or (right-left+1)*(bottom-top+1) > 100000 or bottom > 1048576 or right > 16384:
                    raise ValueError('Formatting range must contain at most 100000 cells within XLSX limits')
                for row in sheet.iter_rows(min_row=top, max_row=bottom, min_col=left, max_col=right):
                    for cell in row:
                        if 'number_format' in fmt:
                            cell.number_format = fmt['number_format']
                        font = copy(cell.font)
                        if 'bold' in fmt: font.bold = fmt['bold']
                        if 'font_color' in fmt: font.color = fmt['font_color']
                        if 'font_size' in fmt: font.sz = fmt['font_size']
                        cell.font = font
                        if 'fill' in fmt:
                            cell.fill = openpyxl.styles.PatternFill('solid', fgColor=fmt['fill'])
                        if 'alignment' in fmt or 'wrap_text' in fmt:
                            alignment = copy(cell.alignment)
                            if 'alignment' in fmt: alignment.horizontal = fmt['alignment']
                            if 'wrap_text' in fmt: alignment.wrap_text = fmt['wrap_text']
                            cell.alignment = alignment
            if 'freeze_panes' in change: sheet.freeze_panes = change['freeze_panes']
            if 'auto_filter' in change: sheet.auto_filter.ref = change['auto_filter']
            for column, width in change.get('column_widths', {}).items():
                sheet.column_dimensions[column].width = width
        book.calculation = openpyxl.workbook.properties.CalcProperties(fullCalcOnLoad=True)
        book.save(target)
    finally:
        book.close()
    check = openpyxl.load_workbook(target); check.close()
    return {'formulas_recalculated': False, 'calculation_notice': 'Formulas are preserved; Excel or LibreOffice must recalculate them.'}


def create_slides(target, args):
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import MSO_AUTO_SIZE
    deck = Presentation(); deck.slide_width = Inches(13.333); deck.slide_height = Inches(7.5)
    theme = {'background': 'FFFFFF', 'foreground': '243247', 'accent': 'C76D4D', 'font': 'Aptos'}
    theme.update(args.get('theme', {}))
    def textbox(slide, x, y, width, height, text, size, color, bold=False):
        shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(width), Inches(height))
        frame = shape.text_frame; frame.word_wrap = True; frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
        frame.text = text
        for paragraph in frame.paragraphs:
            paragraph.font.name = theme['font']; paragraph.font.size = Pt(size)
            paragraph.font.bold = bold; paragraph.font.color.rgb = RGBColor.from_string(color)
        return shape
    for source in args['slides']:
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.background.fill.solid(); slide.background.fill.fore_color.rgb = RGBColor.from_string(theme['background'])
        layout = source.get('layout', 'content')
        title_y = 2.2 if layout == 'section' else .5
        textbox(slide, .7, title_y, 11.9, 1.15, source['title'], 38 if layout == 'section' else 32, theme['accent'], True)
        if source.get('body'):
            textbox(slide, .75, 3.5 if layout == 'section' else 1.8, 11.8, 1.4 if source.get('bullets') or source.get('table') else 4.8,
                    source['body'], 24, theme['foreground'])
        if source.get('bullets'):
            textbox(slide, .9, 3.3 if source.get('body') else 1.9, 11.4, 3.4 if source.get('body') else 4.8,
                    '\n'.join('\u2022 ' + bullet for bullet in source['bullets']), 24, theme['foreground'])
        table_source = source.get('table')
        if table_source:
            rows = [table_source['headers'], *table_source['rows']]
            count = len(table_source['headers'])
            if any(len(row) != count for row in rows):
                raise ValueError('Slide table rows must match the header column count')
            table = slide.shapes.add_table(len(rows), count, Inches(.75), Inches(3.3 if source.get('body') else 1.9),
                                            Inches(11.8), Inches(min(4.5, .45 * len(rows)))).table
            for row_index, row in enumerate(rows):
                for column_index, value in enumerate(row):
                    cell = table.cell(row_index, column_index); cell.text = str(value)
                    cell.fill.solid(); cell.fill.fore_color.rgb = RGBColor.from_string(theme['accent'] if row_index == 0 else theme['background'])
                    for paragraph in cell.text_frame.paragraphs:
                        paragraph.font.name = theme['font']; paragraph.font.size = Pt(16)
                        paragraph.font.bold = row_index == 0
                        paragraph.font.color.rgb = RGBColor.from_string(theme['background'] if row_index == 0 else theme['foreground'])
        if source.get('notes'):
            slide.notes_slide.notes_text_frame.text = source['notes']
    deck.save(target); Presentation(target)


def inspect_slides(path):
    from pptx import Presentation
    deck = Presentation(path)
    return {'slide_count': len(deck.slides), 'slides': [
        {'slide': index + 1, 'shapes': [
            {'shape_id': shape.shape_id, 'name': shape.name,
             'text': shape.text[:6000] if shape.has_text_frame else None,
             'table': [[cell.text[:1000] for cell in row.cells] for row in shape.table.rows][:20] if shape.has_table else None}
            for shape in list(slide.shapes)[:100]],
         'notes': slide.notes_slide.notes_text_frame.text[:6000] if slide.has_notes_slide else ''}
        for index, slide in enumerate(islice(deck.slides, 100))],
        'truncated': len(deck.slides) > 100}


def edit_slides(source, target, args):
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    deck = Presentation(source)
    def slide_at(number):
        if type(number) is not int or not 1 <= number <= len(deck.slides):
            raise ValueError('Slide number must refer to an existing slide (1-based)')
        return deck.slides[number - 1]
    for change in args.get('changes', []):
        slide = slide_at(change['slide'])
        shape = next((item for item in slide.shapes if item.shape_id == change['shape_id']), None)
        if shape is None or not shape.has_text_frame:
            raise ValueError('Text shape was not found; inspect slides first')
        shape.text_frame.text = change['text']
        for paragraph in shape.text_frame.paragraphs:
            if 'font_size' in change: paragraph.font.size = Pt(change['font_size'])
            if 'color' in change: paragraph.font.color.rgb = RGBColor.from_string(change['color'])
            if 'bold' in change: paragraph.font.bold = change['bold']
    for note in args.get('notes', []):
        slide_at(note['slide']).notes_slide.notes_text_frame.text = note['text']
    if 'order' in args:
        order = args['order']
        if not order or len(order) != len(set(order)) or any(type(number) is not int or not 1 <= number <= len(deck.slides) for number in order):
            raise ValueError('order must contain unique existing slide numbers')
        ids = list(deck.slides._sldIdLst)
        kept = set(order)
        for index, item in enumerate(ids, 1):
            deck.slides._sldIdLst.remove(item)
            if index not in kept: deck.part.drop_rel(item.rId)
        for number in order: deck.slides._sldIdLst.append(ids[number - 1])
    for item in args.get('add_slides', []):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        for text, y, height, size in [(item['title'], .5, 1.2, 32), (item.get('body', ''), 1.9, 4.7, 24)]:
            box = slide.shapes.add_textbox(Inches(.7), Inches(y), deck.slide_width - Inches(1.4), Inches(height))
            box.text_frame.word_wrap = True; box.text_frame.text = text
            for paragraph in box.text_frame.paragraphs: paragraph.font.size = Pt(size)
    if len(deck.slides) > 200: raise ValueError('Edited presentation exceeds 200 slides')
    deck.save(target)
    check = Presentation(target)
    return {'slide_count': len(check.slides), 'formatting_notice': 'Only selected shape text and explicit style properties were edited. Inspect previews for text fit.'}


def chart_spreadsheet(source, target, args):
    import openpyxl
    from openpyxl.chart import BarChart, LineChart, PieChart, Reference
    from openpyxl.utils.cell import range_boundaries
    book = openpyxl.load_workbook(source)
    try:
        if args['sheet'] not in book.sheetnames: raise ValueError('Requested sheet was not found')
        sheet = book[args['sheet']]
        left, top, right, bottom = range_boundaries(args['range'])
        if right <= left or bottom <= top or (right-left+1)*(bottom-top+1) > 100000:
            raise ValueError('Chart range needs a header row, category column and numeric data, with at most 100000 cells')
        chart = {'bar': BarChart, 'line': LineChart, 'pie': PieChart}[args['chart_type']]()
        chart.title = args.get('title', 'Chart')
        end_column = left + 1 if args['chart_type'] == 'pie' else right
        chart.add_data(Reference(sheet, min_col=left+1, max_col=end_column, min_row=top, max_row=bottom), titles_from_data=True)
        chart.set_categories(Reference(sheet, min_col=left, min_row=top+1, max_row=bottom))
        sheet.add_chart(chart, args.get('anchor', 'E2'))
        book.save(target)
    finally:
        book.close()
    check = openpyxl.load_workbook(target); check.close()
    return {'chart_type': args['chart_type'], 'sheet': args['sheet']}
