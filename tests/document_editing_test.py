"""Run with: python -m unittest discover -s tests -p 'document_editing_test.py'."""
import sys
import tempfile
import unittest
from pathlib import Path

from pypdf import PdfReader
from reportlab.pdfgen.canvas import Canvas
from openpyxl import Workbook, load_workbook
from pptx import Presentation

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'runtime' / 'worker'))
import document_editing as editing


class DocumentEditingTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)

    def tearDown(self):
        self.temp.cleanup()

    def test_pdf_select_reorders_rotates_and_adds_text_without_changing_source(self):
        self.assertTrue(hasattr(editing, 'edit_pdf'), 'PDF editing is not implemented')
        source = self.root / 'source.pdf'
        canvas = Canvas(str(source))
        for label in ['First page', 'Second page', 'Third page']:
            canvas.drawString(40, 700, label)
            canvas.showPage()
        canvas.save()
        target = self.root / 'edited.pdf'
        editing.edit_pdf(source, target, {
            'pages': [3, 1], 'rotations': [{'page': 1, 'degrees': 90}],
            'overlays': [{'page': 2, 'text': 'Reviewed', 'x': 40, 'y': 40, 'font_size': 16}],
        })
        actual = PdfReader(target)
        self.assertEqual(len(actual.pages), 2)
        self.assertIn('Third page', actual.pages[0].extract_text())
        self.assertEqual(actual.pages[0].rotation, 90)
        self.assertIn('Reviewed', actual.pages[1].extract_text())
        self.assertEqual(len(PdfReader(source).pages), 3)
        info = editing.inspect_pdf(target)
        self.assertEqual(info['page_count'], 2)
        self.assertEqual(info['pages'][0]['rotation'], 90)

    def test_pdf_invalid_page_does_not_write_output(self):
        self.assertTrue(hasattr(editing, 'edit_pdf'), 'PDF editing is not implemented')
        source = self.root / 'source.pdf'
        canvas = Canvas(str(source)); canvas.drawString(10, 10, 'Page'); canvas.save()
        target = self.root / 'bad.pdf'
        with self.assertRaisesRegex(ValueError, 'page'):
            editing.edit_pdf(source, target, {'pages': [2]})
        self.assertFalse(target.exists())

    def test_pdf_duplicate_selected_pages_can_be_edited_independently(self):
        source = self.root / 'source.pdf'
        canvas = Canvas(str(source)); canvas.drawString(10, 10, 'Original'); canvas.save()
        target = self.root / 'duplicate.pdf'
        editing.edit_pdf(source, target, {'pages': [1, 1], 'overlays': [
            {'page': 2, 'text': 'Copy only', 'x': 10, 'y': 50},
        ]})
        actual = PdfReader(target)
        self.assertNotIn('Copy only', actual.pages[0].extract_text())
        self.assertIn('Copy only', actual.pages[1].extract_text())

    def test_spreadsheet_inspection_bounds_columns_and_rows(self):
        source = self.root / 'sample.xlsx'; book = Workbook()
        book.active.append([1, 2, 3]); book.active.append([4, 5, 6]); book.save(source); book.close()
        info = editing.inspect_spreadsheet(source, {'max_rows': 1, 'max_columns': 2})
        self.assertEqual(info['sheets'][0]['rows'], [[1, 2]])
        self.assertTrue(info['truncated'])

    def test_xlsx_reversed_format_range_is_rejected_without_output(self):
        source = self.root / 'sample.xlsx'; book = Workbook(); book.save(source); book.close()
        target = self.root / 'bad.xlsx'
        with self.assertRaisesRegex(ValueError, 'range'):
            editing.edit_spreadsheet(source, target, {'sheets': [{'name': 'Sheet', 'formats': [
                {'range': 'B3:A1', 'bold': True},
            ]}]})
        self.assertFalse(target.exists())

    def test_pdf_fill_preserves_canonical_field_and_widget_appearance(self):
        self.assertTrue(hasattr(editing, 'edit_pdf'), 'PDF editing is not implemented')
        source = self.root / 'form.pdf'
        canvas = Canvas(str(source))
        canvas.acroForm.textfield(name='customer', x=40, y=600, width=200, height=30)
        canvas.showPage(); canvas.save()
        target = self.root / 'filled.pdf'
        editing.edit_pdf(source, target, {'form_fields': {'customer': 'Ada'}})
        reader = PdfReader(target)
        self.assertEqual(reader.get_fields()['customer']['/V'], 'Ada')
        widget = reader.pages[0]['/Annots'][0].get_object()
        self.assertTrue(widget['/AP']['/N'])
        self.assertEqual(PdfReader(source).get_fields()['customer']['/V'], '')

    def test_xlsx_edit_retains_sheets_formulas_and_formats(self):
        self.assertTrue(hasattr(editing, 'edit_spreadsheet'), 'Spreadsheet editing is not implemented')
        source = self.root / 'source.xlsx'
        book = Workbook(); sheet = book.active; sheet.title = 'Sales'
        sheet.append(['Amount', 'Total']); sheet.append([25, '=A2*2'])
        book.create_sheet('Untouched')['A1'] = 'Keep me'
        book.save(source); book.close()
        target = self.root / 'edited.xlsx'
        editing.edit_spreadsheet(source, target, {'sheets': [{
            'name': 'Sales', 'cells': [{'cell': 'A2', 'value': 40}, {'cell': 'B3', 'value': '=SUM(B2:B2)'}],
            'formats': [{'range': 'A2:B3', 'number_format': '#,##0.00', 'bold': True, 'fill': 'FFF0CC'}],
            'freeze_panes': 'B2', 'auto_filter': 'A1:B3', 'column_widths': {'A': 24},
        }]})
        actual = load_workbook(target)
        self.assertEqual(actual['Sales']['A2'].value, 40)
        self.assertEqual(actual['Sales']['B2'].value, '=A2*2')
        self.assertEqual(actual['Sales']['B3'].value, '=SUM(B2:B2)')
        self.assertEqual(actual['Sales']['A2'].number_format, '#,##0.00')
        self.assertTrue(actual['Sales']['A2'].font.bold)
        self.assertEqual(actual['Sales'].freeze_panes, 'B2')
        self.assertEqual(actual['Sales'].auto_filter.ref, 'A1:B3')
        self.assertEqual(actual['Untouched']['A1'].value, 'Keep me')
        actual.close()
        self.assertEqual(load_workbook(source)['Sales']['A2'].value, 25)

    def test_spreadsheet_inspect_reports_actual_truncation_and_missing_sheet(self):
        self.assertTrue(hasattr(editing, 'inspect_spreadsheet'), 'Bounded spreadsheet inspection is not implemented')
        source = self.root / 'sample.xlsx'; book = Workbook()
        book.active.append(['One', '=1+2']); book.save(source); book.close()
        info = editing.inspect_spreadsheet(source, {})
        self.assertFalse(info['truncated'])
        self.assertEqual(info['sheets'][0]['rows'], [['One', '=1+2']])
        with self.assertRaisesRegex(ValueError, 'sheet'):
            editing.inspect_spreadsheet(source, {'sheet': 'Missing'})

    def test_slides_theme_bullets_table_and_notes_survive_reopen(self):
        self.assertTrue(hasattr(editing, 'create_slides'), 'Rich slide creation is not implemented')
        target = self.root / 'deck.pptx'
        editing.create_slides(target, {'theme': {'background': '112233', 'foreground': 'FFFFFF', 'accent': 'FFAA00'}, 'slides': [
            {'title': 'Roadmap', 'layout': 'section', 'body': 'Next quarter', 'notes': 'Speaker context'},
            {'title': 'Priorities', 'bullets': ['Ship', 'Learn']},
            {'title': 'Plan', 'table': {'headers': ['Phase', 'Owner'], 'rows': [['Build', 'Ada']]}},
        ]})
        deck = Presentation(target)
        self.assertEqual(len(deck.slides), 3)
        self.assertEqual(str(deck.slides[0].background.fill.fore_color.rgb), '112233')
        self.assertIn('Speaker context', deck.slides[0].notes_slide.notes_text_frame.text)
        self.assertTrue(any('Ship' in getattr(shape, 'text', '') for shape in deck.slides[1].shapes))
        table = next(shape.table for shape in deck.slides[2].shapes if shape.has_table)
        self.assertEqual(table.cell(1, 1).text, 'Ada')


if __name__ == '__main__':
    unittest.main()
